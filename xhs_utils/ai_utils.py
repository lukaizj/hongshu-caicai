import json
import os
import threading
import time
from datetime import datetime

import openpyxl

PROJECT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENV_PATH = os.path.join(PROJECT_DIR, ".env")
EXCEL_DIR = os.path.join(PROJECT_DIR, "datas", "excel_datas")

AI_RESULTS = {}
AI_RESULTS_LOCK = threading.Lock()
AI_RESULTS_TTL = 7200
MAX_AI_RESULTS = 30

# ── Skill templates ─────────────────────────────────────────────

AI_SKILLS = {
    "trend": {
        "name": "内容趋势分析",
        "description": "分析笔记内容趋势，提取热门话题和模式",
        "system_prompt": """你是一位专业的小红书内容分析师。请基于以下笔记数据，进行内容趋势分析：

## 分析维度
1. **热门主题**：哪些关键词、标签或话题出现频率最高
2. **内容形式偏好**：图集 vs 视频的占比和互动表现
3. **互动规律**：点赞、收藏、评论、分享之间的关联模式
4. **时间趋势**：发布时间的分布和对互动量的影响
5. **IP 地域分布**：主要发布地区及其内容特征

请用中文回答，结构清晰，使用标题和列表。如果数据量充足，给出具体的数字支撑。""",
    },
    "sentiment": {
        "name": "评论情感分析",
        "description": "分析评论的情感倾向和用户反馈",
        "system_prompt": """你是一位专业的用户研究分析师。请基于以下小红书评论数据，进行情感分析：

## 分析维度
1. **情感分布**：评论的正面、负面、中性比例
2. **高频关键词**：评论中出现最多的词汇和短语
3. **关注焦点**：用户最关心的产品/内容维度是什么
4. **改进建议**：负面评论中反映的改进方向
5. **口碑亮点**：正面评论中提及的核心卖点

请用中文回答，结构清晰。如果数据允许，按情感分类列出典型评论片段。""",
    },
    "strategy": {
        "name": "选题策略建议",
        "description": "根据数据提供内容选题和创作策略",
        "system_prompt": """你是一位资深的小红书内容策略顾问。请基于以下笔记和评论数据，提供选题策略建议：

## 分析维度
1. **爆款基因**：高互动笔记的共同特征（标题风格、内容类型、发布时间等）
2. **受众洞察**：评论反映出的用户画像和核心需求
3. **内容缺口**：该领域尚未被充分满足的用户需求
4. **选题方向**：建议的 3-5 个具体选题方向
5. **差异化策略**：如何在同质化内容中脱颖而出

请用中文回答，建议具体可执行，每个方向附带简要理由。""",
    },
    "persona": {
        "name": "用户画像分析",
        "description": "分析目标用户的特征和行为模式",
        "system_prompt": """你是一位数据分析师，负责构建用户画像。请基于以下数据，分析目标用户群体：

## 分析维度
1. **用户身份**：根据昵称、评论内容等推测用户身份特征
2. **消费行为**：评论中的消费决策模式、价格敏感度、购买动机
3. **内容偏好**：用户最常关注和互动的内容类型
4. **使用场景**：评论中反映出的产品使用场景和痛点
5. **决策旅程**：从注意到购买的典型决策路径

请用中文回答，给出清晰的用户画像描述。""",
    },
    "viral": {
        "name": "爆款特征提取",
        "description": "提取高互动笔记的共同特征和爆款规律",
        "system_prompt": """你是一位内容增长分析师。请基于以下笔记数据，提取爆款笔记的特征规律：

## 分析维度
1. **标题分析**：高互动笔记的标题长度、关键词、句式特征
2. **内容特征**：描述文案的篇幅、结构、情绪基调
3. **互动比率**：点赞/收藏/评论比率判断内容类型（干货 vs. 娱乐 vs. 种草）
4. **形式因素**：图集 vs 视频、标签策略对互动的影响
5. **时间因素**：发布时间、时段与互动量的关联

请用中文回答，用具体数据说明每个特征。""",
    },
}

PROVIDER_LABELS = {
    "claude": "Claude (Anthropic)",
    "openai": "OpenAI",
    "deepseek": "DeepSeek",
}

PROVIDER_DEFAULTS = {
    "claude": "claude-sonnet-4-6",
    "openai": "gpt-4o",
    "deepseek": "deepseek-chat",
}

# ── Config management ────────────────────────────────────────────

def _read_env_lines():
    if os.path.exists(ENV_PATH):
        with open(ENV_PATH, "r", encoding="utf-8") as f:
            return f.read().splitlines()
    return []


def _write_env_lines(lines):
    with open(ENV_PATH, "w", encoding="utf-8") as f:
        f.write("\n".join(lines).rstrip() + "\n")


def _quote_env_value(value):
    return "'" + value.replace("\\", "\\\\").replace("'", "\\'") + "'"


def load_ai_config():
    """返回 {configured, provider, model, enabled, key_masked, provider_label}"""
    provider = os.environ.get("AI_PROVIDER", "").strip()
    api_key = os.environ.get("AI_API_KEY", "").strip()
    model = os.environ.get("AI_MODEL", "").strip()
    enabled = os.environ.get("AI_ENABLED", "").strip() == "1"

    if api_key:
        key_masked = api_key[:4] + "***" + api_key[-4:] if len(api_key) > 8 else "***"
    else:
        key_masked = ""

    return {
        "configured": bool(provider and api_key),
        "provider": provider,
        "model": model,
        "enabled": enabled,
        "key_masked": key_masked,
        "provider_label": PROVIDER_LABELS.get(provider, provider),
    }


def save_ai_config(provider, api_key, model, enabled):
    provider = (provider or "").strip()
    api_key = (api_key or "").strip()
    model = (model or "").strip()
    if not provider or not api_key:
        raise ValueError("提供商和 API Key 不能为空")

    enabled_val = "1" if enabled else "0"
    if not model:
        model = PROVIDER_DEFAULTS.get(provider, "")

    lines = _read_env_lines()
    updates = {
        "AI_PROVIDER": provider,
        "AI_API_KEY": api_key,
        "AI_MODEL": model,
        "AI_ENABLED": enabled_val,
    }
    written = {k: False for k in updates}

    new_lines = []
    for line in lines:
        matched = False
        for key, val in updates.items():
            if line.startswith(key + "="):
                new_lines.append(key + "=" + _quote_env_value(val))
                written[key] = True
                matched = True
                break
        if not matched:
            new_lines.append(line)

    for key, val in updates.items():
        if not written[key]:
            new_lines.append(key + "=" + _quote_env_value(val))

    _write_env_lines(new_lines)

    os.environ["AI_PROVIDER"] = provider
    os.environ["AI_API_KEY"] = api_key
    os.environ["AI_MODEL"] = model
    os.environ["AI_ENABLED"] = enabled_val

    return load_ai_config()


# ── Excel → text conversion ──────────────────────────────────────

def excel_to_text(excel_path, max_rows=50):
    """将 Excel 文件转换为结构化文本供 LLM 分析"""
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    total_rows = ws.max_row - 1

    if total_rows == 0:
        return "（该文件无数据行）"

    is_comments = "评论内容" in headers
    lines = [
        f"## 文件：{os.path.basename(excel_path)}",
        f"- 总行数：{total_rows}",
        f"- 列名：{' | '.join(headers)}",
        "",
    ]

    rows_to_show = min(total_rows, max_rows)
    lines.append(f"### 数据行（前 {rows_to_show} 行）")

    for row_idx in range(2, 2 + rows_to_show):
        row_data = {}
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            row_data[header] = str(cell.value) if cell.value is not None else ""

        if is_comments:
            compact = {
                "昵称": row_data.get("昵称", ""),
                "评论内容": row_data.get("评论内容", ""),
                "点赞": row_data.get("点赞数量", "0"),
                "时间": row_data.get("上传时间", ""),
                "ip": row_data.get("ip归属地", ""),
            }
        else:
            compact = {
                "标题": row_data.get("标题", ""),
                "描述": row_data.get("描述", ""),
                "点赞": row_data.get("点赞数量", "0"),
                "收藏": row_data.get("收藏数量", "0"),
                "评论": row_data.get("评论数量", "0"),
                "分享": row_data.get("分享数量", "0"),
                "标签": row_data.get("标签", ""),
                "时间": row_data.get("上传时间", ""),
                "ip": row_data.get("ip归属地", ""),
                "昵称": row_data.get("昵称", ""),
                "类型": row_data.get("笔记类型", ""),
            }
        lines.append(json.dumps(compact, ensure_ascii=False))

    if total_rows > max_rows:
        lines.append(f"\n（共 {total_rows} 行，仅展示前 {max_rows} 行）")

    wb.close()
    return "\n".join(lines)


def read_excel_files_for_job(job):
    """从 job 中读取关联的 Excel 文件内容"""
    notes_text = ""
    comments_text = ""
    downloads = job.get("downloads", [])
    for dl in downloads:
        filename = dl.get("file", "")
        filepath = os.path.join(EXCEL_DIR, filename)
        if not os.path.exists(filepath):
            continue
        if "_comments" in filename:
            comments_text = excel_to_text(filepath)
        else:
            notes_text = excel_to_text(filepath)
    return notes_text, comments_text


# ── Prompt building ──────────────────────────────────────────────

def build_prompt(skill_type, custom_prompt, notes_text, comments_text):
    """返回 (system_prompt, user_message)"""
    if skill_type == "custom":
        system_prompt = custom_prompt or "请分析以下小红书数据。"
    else:
        skill = AI_SKILLS.get(skill_type, AI_SKILLS["trend"])
        system_prompt = skill["system_prompt"]
        if custom_prompt:
            system_prompt += "\n\n## 额外要求\n" + custom_prompt

    parts = []
    if notes_text:
        parts.append(f"# 笔记数据\n\n{notes_text}")
    if comments_text:
        parts.append(f"# 评论数据\n\n{comments_text}")
    if not parts:
        parts.append("（无可用数据）")

    user_message = "\n\n".join(parts)
    return system_prompt, user_message


# ── LLM calling ──────────────────────────────────────────────────

def _call_claude(api_key, model, system, message):
    import anthropic

    client = anthropic.Anthropic(api_key=api_key)
    resp = client.messages.create(
        model=model,
        system=system,
        messages=[{"role": "user", "content": message}],
        max_tokens=4096,
        temperature=0.7,
    )
    return resp.content[0].text


def _call_openai(api_key, model, system, message, base_url=None):
    from openai import OpenAI

    kwargs = {"api_key": api_key}
    if base_url:
        kwargs["base_url"] = base_url
    client = OpenAI(**kwargs)
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system},
            {"role": "user", "content": message},
        ],
        max_tokens=4096,
        temperature=0.7,
    )
    return resp.choices[0].message.content


def call_llm(provider, api_key, model, system, message):
    if provider == "claude":
        return _call_claude(api_key, model, system, message)
    elif provider == "openai":
        return _call_openai(api_key, model, system, message)
    elif provider == "deepseek":
        return _call_openai(api_key, model, system, message, base_url="https://api.deepseek.com")
    else:
        raise ValueError(f"不支持的 AI 提供商: {provider}")


# ── Background analysis ─────────────────────────────────────────

def store_ai_result(aid, **fields):
    with AI_RESULTS_LOCK:
        AI_RESULTS[aid] = fields
        AI_RESULTS[aid]["updated_at"] = time.time()
        # cleanup old results
        if len(AI_RESULTS) > MAX_AI_RESULTS:
            oldest = sorted(AI_RESULTS.items(), key=lambda x: x[1].get("updated_at", 0))
            for old_id, _ in oldest[: len(oldest) - MAX_AI_RESULTS]:
                del AI_RESULTS[old_id]


def get_ai_result(aid):
    with AI_RESULTS_LOCK:
        r = AI_RESULTS.get(aid)
        if r is None:
            return None
        # check TTL
        if time.time() - r.get("updated_at", 0) > AI_RESULTS_TTL:
            del AI_RESULTS[aid]
            return None
        return dict(r)


def build_draft_prompt(notes_text, comments_text, source_analysis, style_profile, topic_angle, target_audience, tone, draft_count, include_image_prompts):
    system_prompt = """你是一位小红书内容策划和原创改写专家。你的任务是基于爆款样本提炼结构，但必须生成原创内容，不能照抄标题、正文、图片描述或品牌表达。

输出必须使用 Markdown，结构固定为：
# 标题候选
# 正文草稿
# 封面文案
# 配图建议 / 图片生成提示词
# 标签
# 发布前检查
# 原创改写与风险提醒

要求：
- 标题候选数量不少于 5 个。
- 正文要适合小红书发布，包含开头钩子、主体段落、结尾互动引导。
- 标签控制在 8-15 个。
- 不使用“复制”“仿写”“搬运”等表述。
- 风格要贴合用户给定账号定位，不要机械堆关键词。"""
    image_prompt_rule = "请为每篇草稿输出封面图和正文配图的图片生成提示词。" if include_image_prompts else "不需要输出图片生成提示词，只输出配图方向。"
    user_message = f"""# 创作目标

- 账号定位 / 风格描述：{style_profile or '未提供'}
- 切入选题角度：{topic_angle or '请基于样本自行建议'}
- 目标人群：{target_audience or '未提供'}
- 语气：{tone or '自然、有用、有记忆点'}
- 生成篇数：{draft_count}
- 图片要求：{image_prompt_rule}

# 已有爆款拆解

{source_analysis or '（无单独爆款拆解结果，请直接基于样本数据判断。）'}

# 爆款样本笔记数据

{notes_text or '（无笔记数据）'}

# 评论数据

{comments_text or '（无评论数据）'}
"""
    return system_prompt, user_message


def run_draft_generation(draft_id, job, draft_request, ai_config):
    try:
        store_ai_result(draft_id, status="running", result="", error="", type="draft", job_id=job.get("id", ""))
        notes_text, comments_text = read_excel_files_for_job(job)
        if not notes_text and not comments_text:
            raise ValueError("未找到可生成草稿的 Excel 数据，请确认采集已完成")

        system_prompt, user_message = build_draft_prompt(
            notes_text,
            comments_text,
            draft_request.get("source_analysis", ""),
            draft_request.get("style_profile", ""),
            draft_request.get("topic_angle", ""),
            draft_request.get("target_audience", ""),
            draft_request.get("tone", ""),
            draft_request.get("draft_count", 1),
            draft_request.get("include_image_prompts", True),
        )
        result = call_llm(ai_config["provider"], ai_config["api_key"], ai_config["model"], system_prompt, user_message)
        store_ai_result(
            draft_id,
            status="done",
            result=result,
            error="",
            type="draft",
            job_id=job.get("id", ""),
            title="小红书草稿",
            completed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            metadata={
                "tone": draft_request.get("tone", ""),
                "draft_count": draft_request.get("draft_count", 1),
                "include_image_prompts": draft_request.get("include_image_prompts", True),
            },
        )
    except Exception as exc:
        store_ai_result(draft_id, status="error", result="", error=str(exc), type="draft")


def run_analysis(analysis_id, job, skill_type, custom_prompt, provider, api_key, model):
    try:
        store_ai_result(analysis_id, status="running", result="", error="")

        # 1. read excel
        notes_text, comments_text = read_excel_files_for_job(job)
        if not notes_text and not comments_text:
            raise ValueError("未找到可分析的 Excel 数据，请确认采集已完成")

        # 2. build prompt
        system_prompt, user_message = build_prompt(skill_type, custom_prompt, notes_text, comments_text)

        # 3. call LLM
        result = call_llm(provider, api_key, model, system_prompt, user_message)

        # 4. store
        store_ai_result(
            analysis_id,
            status="done",
            result=result,
            error="",
            skill_type=skill_type,
            skill_name=AI_SKILLS.get(skill_type, {}).get("name", "自定义分析"),
            completed_at=datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        )
    except Exception as exc:
        store_ai_result(analysis_id, status="error", result="", error=str(exc))
